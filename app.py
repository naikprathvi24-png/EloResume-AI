
from flask import Flask, render_template, request, redirect, url_for, session, flash, make_response
import sqlite3
import json
from io import BytesIO
from werkzeug.security import generate_password_hash, check_password_hash
from weasyprint import HTML
from flask_mail import Mail, Message

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import pymupdf
import os
from uuid import uuid4
from werkzeug.utils import secure_filename
from dotenv import load_dotenv
from google import genai
import secrets
import hashlib
import time
import hmac

app = Flask(__name__)

load_dotenv()
# Gmail SMTP configuration
app.config["MAIL_SERVER"] = os.getenv(
    "MAIL_SERVER", "smtp.gmail.com"
)
app.config["MAIL_PORT"] = int(
    os.getenv("MAIL_PORT", "587")
)
app.config["MAIL_USE_TLS"] = True
app.config["MAIL_USERNAME"] = os.getenv("MAIL_USERNAME")
app.config["MAIL_PASSWORD"] = os.getenv("MAIL_PASSWORD")
app.config["MAIL_DEFAULT_SENDER"] = os.getenv(
    "MAIL_DEFAULT_SENDER"
)

mail = Mail(app)

# -----------------------------
# OTP EMAIL AND SECURITY
# -----------------------------

def send_otp_email(email, otp, purpose="verification"):
    message = Message(
        subject="Your EloResume AI Verification Code",
        recipients=[email]
    )

    message.body = f"""
Hello,

Your EloResume AI verification code is:

{otp}

This code expires in 5 minutes and can only be used once.

Purpose: {purpose}

If you did not request this code, please ignore this email.

Regards,
EloResume AI
"""

    mail.send(message)


def generate_otp():
    """Generate a secure six-digit OTP."""
    return f"{secrets.randbelow(1000000):06d}"


def hash_otp(otp):
    """Hash an OTP before storing it."""
    return hashlib.sha256(otp.encode()).hexdigest()
def create_otp(email, purpose, user_id=None, pending_data=None):
    otp = generate_otp()
    otp_hash = hash_otp(otp)

    now = int(time.time())
    expires_at = now + 300

    conn = get_db()
    conn.execute(
        """
        DELETE FROM otp_verifications
        WHERE email = ? AND purpose = ?
        """,
        (email, purpose)
    )

    conn.execute(
        """
        INSERT INTO otp_verifications
        (email, purpose, otp_hash, expires_at, attempts,
         created_at, user_id, pending_data)
        VALUES (?, ?, ?, ?, 0, ?, ?, ?)
        """,
        (
            email,
            purpose,
            otp_hash,
            expires_at,
            now,
            user_id,
            json.dumps(pending_data) if pending_data else None
        )
    )

    conn.commit()
    conn.close()

    return otp
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")

gemini_client = None

if GEMINI_API_KEY:
    gemini_client = genai.Client(
        api_key=GEMINI_API_KEY
    )
def generate_ai_text(prompt):

    if gemini_client is None:
        return None

    response = gemini_client.models.generate_content(
        model="gemini-3.5-flash-lite",
        contents=prompt
    )

    return response.text.strip()

# Secret key for login sessions
app.secret_key = "eloresume-secret-key-change-this-later"

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

DATABASE_DIR = os.path.join(BASE_DIR, "database")
os.makedirs(DATABASE_DIR, exist_ok=True)

DATABASE = os.path.join(DATABASE_DIR, "eloresume.db")

# -----------------------------
# DATABASE CONNECTION
# -----------------------------
def get_db():
    connection = sqlite3.connect(DATABASE)
    connection.row_factory = sqlite3.Row
    return connection

# -----------------------------
# CREATE DATABASE TABLE
# -----------------------------

def init_db():

    connection = sqlite3.connect(DATABASE)
    connection.row_factory = sqlite3.Row

    # OTP verification table
    connection.execute("""
        CREATE TABLE IF NOT EXISTS otp_verifications (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            email TEXT NOT NULL,
            purpose TEXT NOT NULL,
            otp_hash TEXT NOT NULL,
            expires_at INTEGER NOT NULL,
            attempts INTEGER DEFAULT 0,
            created_at INTEGER NOT NULL,
            user_id INTEGER,
            pending_data TEXT
        )
    """)

    # Users table
    connection.execute("""
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            email TEXT UNIQUE NOT NULL,
            password TEXT NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    """)
    

    # Resumes table
    connection.execute("""
        CREATE TABLE IF NOT EXISTS resumes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            resume_name TEXT NOT NULL,
            resume_type TEXT NOT NULL,
            style TEXT NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (user_id) REFERENCES users(id)
        )
    """)

    # Add resume_data column if it does not already exist
    columns = connection.execute(
        "PRAGMA table_info(resumes)"
    ).fetchall()

    column_names = [column["name"] for column in columns]

    if "resume_data" not in column_names:
        connection.execute(
            "ALTER TABLE resumes ADD COLUMN resume_data TEXT"
        )

    connection.commit()
    connection.close()
    
    
@app.route("/api/ai/generate", methods=["POST"])
def ai_generate():

    if "user_id" not in session:
        return {
            "success": False,
            "message": "Please login first."
        }, 401

    data = request.get_json()

    if not data:
        return {
            "success": False,
            "message": "No request data received."
        }, 400

    prompt = data.get("prompt", "").strip()

    if not prompt:
        return {
            "success": False,
            "message": "Please provide a prompt."
        }, 400

    try:

        result = generate_ai_text(prompt)

        if not result:
            return {
                "success": False,
                "message": "Gemini API is not configured."
            }, 500

        return {
            "success": True,
            "result": result
        }

    except Exception as error:

        print("Gemini Error:", error)

        return {
            "success": False,
            "message": "Unable to generate AI content."
        }, 500   
@app.route("/api/ai/analyze-resume", methods=["POST"])
def analyze_resume_with_ai():

    if "user_id" not in session:
        return {
            "success": False,
            "message": "Please login first."
        }, 401

    data = request.get_json()

    if not data:
        return {
            "success": False,
            "message": "No resume data received."
        }, 400

    resume_text = data.get("resume_text", "").strip()

    if not resume_text:
        return {
            "success": False,
            "message": "Resume information is empty."
        }, 400

    prompt = f"""
Analyze the following resume content.

RESUME:
{resume_text}

Provide a professional resume review.

Include these sections:

1. Strengths
2. Missing or incomplete information
3. Skills improvement suggestions
4. Project improvement suggestions
5. Professional wording suggestions
6. General resume improvement tips

Requirements:
- Be clear and practical.
- Do not invent information about the candidate.
- Do not rewrite the entire resume.
- Keep the suggestions suitable for a job application.
- Use simple professional language.
"""

    try:

        result = generate_ai_text(prompt)

        if not result:
            return {
                "success": False,
                "message": "Gemini API is not configured."
            }, 500

        return {
            "success": True,
            "result": result
        }

    except Exception as error:

        print("Resume Analysis Error:", error)

        return {
            "success": False,
            "message": "Unable to analyze resume."
        }, 500        


# -----------------------------
# HOME / LOGIN
# -----------------------------

@app.route("/")
def home():

    if "user_id" in session:
        return redirect(url_for("dashboard"))

    return render_template("login.html")


# -----------------------------
# REGISTER PAGE WITH EMAIL OTP
# -----------------------------

@app.route("/register", methods=["GET", "POST"])
def register():

    if request.method == "POST":

        name = request.form["name"].strip()
        email = request.form["email"].strip().lower()
        password = request.form["password"]
        confirm_password = request.form["confirm_password"]

        if not name or not email or not password:
            flash("Please fill all required fields.", "error")
            return redirect(url_for("register"))

        if password != confirm_password:
            flash("Passwords do not match.", "error")
            return redirect(url_for("register"))

        if len(password) < 6:
            flash("Password must contain at least 6 characters.", "error")
            return redirect(url_for("register"))

        connection = get_db()

        existing_user = connection.execute(
            "SELECT id FROM users WHERE email = ?",
            (email,)
        ).fetchone()

        connection.close()

        if existing_user:
            flash(
                "An account with this email already exists.",
                "error"
            )
            return redirect(url_for("register"))

        hashed_password = generate_password_hash(password)

        pending_data = {
            "name": name,
            "password": hashed_password
        }

        try:
            otp = create_otp(
                email=email,
                purpose="signup",
                pending_data=pending_data
            )

            send_otp_email(
                email,
                otp,
                purpose="signup"
            )

            session["otp_email"] = email
            session["otp_purpose"] = "signup"

            flash(
                "Verification code sent to your email.",
                "success"
            )

            return redirect(url_for("verify_otp"))

        except Exception as error:
            print("OTP email error:", error)

            flash(
                "Unable to send verification email. "
                "Please check your email settings and try again.",
                "error"
            )

            return redirect(url_for("register"))

    return render_template("register.html")
# -----------------------------
# EMAIL OTP VERIFICATION
# -----------------------------


@app.route("/verify-otp", methods=["GET", "POST"])
def verify_otp():

    email = session.get("otp_email")
    purpose = session.get("otp_purpose")

    if not email or purpose not in ("signup", "login"):
        flash("Please start the verification process again.", "error")
        return redirect(url_for("home"))

    connection = get_db()

    otp_record = connection.execute(
        """
        SELECT * FROM otp_verifications
        WHERE email = ? AND purpose = ?
        """,
        (email, purpose)
    ).fetchone()

    if not otp_record:
        connection.close()
        flash("Verification code not found. Please try again.", "error")
        return redirect(url_for("home"))

    if request.method == "POST":

        entered_otp = request.form.get("otp", "").strip()

        if not entered_otp:
            connection.close()
            flash("Please enter the verification code.", "error")
            return redirect(url_for("verify_otp"))

        if int(time.time()) > otp_record["expires_at"]:
            connection.execute(
                "DELETE FROM otp_verifications WHERE id = ?",
                (otp_record["id"],)
            )
            connection.commit()
            connection.close()

            session.pop("otp_email", None)
            session.pop("otp_purpose", None)

            flash("Verification code expired. Please try again.", "error")
            return redirect(url_for("home"))

        if otp_record["attempts"] >= 5:
            connection.close()
            flash("Too many incorrect attempts. Please try again later.", "error")
            return redirect(url_for("home"))

        entered_hash = hash_otp(entered_otp)

        if not hmac.compare_digest(
            entered_hash,
            otp_record["otp_hash"]
        ):
            connection.execute(
                """
                UPDATE otp_verifications
                SET attempts = attempts + 1
                WHERE id = ?
                """,
                (otp_record["id"],)
            )
            connection.commit()
            connection.close()

            flash("Incorrect verification code. Please try again.", "error")
            return redirect(url_for("verify_otp"))

        if purpose == "signup":

            pending_data = json.loads(otp_record["pending_data"])

            connection.execute(
                """
                INSERT INTO users (name, email, password)
                VALUES (?, ?, ?)
                """,
                (
                    pending_data["name"],
                    email,
                    pending_data["password"]
                )
            )

            connection.execute(
                "DELETE FROM otp_verifications WHERE id = ?",
                (otp_record["id"],)
            )

            connection.commit()
            connection.close()

            session.pop("otp_email", None)
            session.pop("otp_purpose", None)

            flash(
                "Email verified! Account created successfully. Please login.",
                "success"
            )

            return redirect(url_for("home"))

        elif purpose == "login":

            user = connection.execute(
                "SELECT * FROM users WHERE id = ?",
                (otp_record["user_id"],)
            ).fetchone()

            if not user:
                connection.close()
                flash("Account not found. Please login again.", "error")
                return redirect(url_for("home"))

            connection.execute(
                "DELETE FROM otp_verifications WHERE id = ?",
                (otp_record["id"],)
            )

            connection.commit()
            connection.close()

            session.pop("otp_email", None)
            session.pop("otp_purpose", None)

            session["user_id"] = user["id"]
            session["user_name"] = user["name"]
            session["user_email"] = user["email"]

            flash("Login successful!", "success")
            return redirect(url_for("dashboard"))

    connection.close()

    return render_template("verify_otp.html", email=email)
# -----------------------------
# LOGIN WITH EMAIL OTP
# -----------------------------

@app.route("/login", methods=["POST"])
def login():

    email = request.form["email"].strip().lower()
    password = request.form["password"]

    connection = get_db()

    user = connection.execute(
        "SELECT * FROM users WHERE email = ?",
        (email,)
    ).fetchone()

    connection.close()

    if user and check_password_hash(user["password"], password):

        try:
            otp = create_otp(
                email=email,
                purpose="login",
                user_id=user["id"]
            )

            send_otp_email(
                email,
                otp,
                purpose="login"
            )

            session["otp_email"] = email
            session["otp_purpose"] = "login"

            flash(
                "Login verification code sent to your email.",
                "success"
            )

            return redirect(url_for("verify_otp"))

        except Exception as error:
            print("Login OTP email error:", error)

            flash(
                "Unable to send verification email. Please try again.",
                "error"
            )

            return redirect(url_for("home"))

    flash("Invalid email or password.", "error")

    return redirect(url_for("home"))
# -----------------------------
# FORGOT PASSWORD
# -----------------------------

@app.route("/forgot-password", methods=["GET", "POST"])
def forgot_password():

    if request.method == "POST":

        email = request.form.get("email", "").strip().lower()

        if not email:
            flash("Please enter your email address.", "error")
            return redirect(url_for("forgot_password"))

        connection = get_db()

        user = connection.execute(
            "SELECT id FROM users WHERE email = ?",
            (email,)
        ).fetchone()

        connection.close()

        if not user:
            flash("No account found with this email address.", "error")
            return redirect(url_for("forgot_password"))

        try:
            otp = create_otp(
                email=email,
                purpose="reset",
                user_id=user["id"]
            )

            send_otp_email(
                email,
                otp,
                purpose="password reset"
            )

            session["otp_email"] = email
            session["otp_purpose"] = "reset"

            flash("Password reset code sent to your email.", "success")

            return redirect(url_for("reset_password"))

        except Exception as error:
            print("Password reset email error:", error)

            flash("Unable to send verification email. Please try again.", "error")

            return redirect(url_for("forgot_password"))

    return render_template("forgot_password.html")
# -----------------------------
# RESET PASSWORD
# -----------------------------

@app.route("/reset-password", methods=["GET", "POST"])
def reset_password():

    email = session.get("otp_email")
    purpose = session.get("otp_purpose")

    if not email or purpose != "reset":
        flash("Please request a password reset first.", "error")
        return redirect(url_for("forgot_password"))

    if request.method == "POST":

        otp = request.form.get("otp", "").strip()
        new_password = request.form.get("new_password", "")
        confirm_password = request.form.get("confirm_password", "")

        if not otp or not new_password or not confirm_password:
            flash("Please fill in all fields.", "error")
            return redirect(url_for("reset_password"))

        if new_password != confirm_password:
            flash("Passwords do not match.", "error")
            return redirect(url_for("reset_password"))

        if len(new_password) < 6:
            flash("Password must contain at least 6 characters.", "error")
            return redirect(url_for("reset_password"))

        connection = get_db()

        record = connection.execute(
            """
            SELECT * FROM otp_verifications
            WHERE email = ? AND purpose = ?
            """,
            (email, "reset")
        ).fetchone()

        if not record:
            connection.close()
            flash("Verification code not found. Request a new code.", "error")
            return redirect(url_for("forgot_password"))

        if int(time.time()) > record["expires_at"]:
            connection.execute(
                "DELETE FROM otp_verifications WHERE id = ?",
                (record["id"],)
            )
            connection.commit()
            connection.close()

            session.pop("otp_email", None)
            session.pop("otp_purpose", None)

            flash("Verification code expired. Request a new code.", "error")
            return redirect(url_for("forgot_password"))

        if record["attempts"] >= 5:
            connection.close()
            flash("Too many incorrect attempts. Request a new code.", "error")
            return redirect(url_for("forgot_password"))

        entered_hash = hash_otp(otp)

        if not hmac.compare_digest(entered_hash, record["otp_hash"]):
            connection.execute(
                """
                UPDATE otp_verifications
                SET attempts = attempts + 1
                WHERE id = ?
                """,
                (record["id"],)
            )
            connection.commit()
            connection.close()

            flash("Incorrect verification code.", "error")
            return redirect(url_for("reset_password"))

        hashed_password = generate_password_hash(new_password)

        connection.execute(
            "UPDATE users SET password = ? WHERE email = ?",
            (hashed_password, email)
        )

        connection.execute(
            "DELETE FROM otp_verifications WHERE id = ?",
            (record["id"],)
        )

        connection.commit()
        connection.close()

        session.pop("otp_email", None)
        session.pop("otp_purpose", None)

        flash("Password reset successfully. Please login.", "success")
        return redirect(url_for("home"))

    return render_template("reset_password.html")

# -----------------------------
# DASHBOARD
# -----------------------------

@app.route("/dashboard")
def dashboard():

    if "user_id" not in session:
        return redirect(url_for("home"))

    connection = get_db()

    resumes = connection.execute(
        """
        SELECT *
        FROM resumes
        WHERE user_id = ?
        ORDER BY updated_at DESC
        """,
        (session["user_id"],)
    ).fetchall()

    connection.close()

    return render_template(
        "dashboard.html",
        user_name=session["user_name"],
        user_email=session["user_email"],
        resumes=resumes
    )
@app.route("/profile", methods=["GET", "POST"])
def profile():

    if "user_id" not in session:
        return redirect(url_for("home"))

    connection = get_db()

    # Create profile photo table if it does not exist
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS profile_photos (
            user_id INTEGER PRIMARY KEY,
            photo_path TEXT
        )
        """
    )

    connection.commit()

    if request.method == "POST":

        name = request.form.get("name", "").strip()
        email = request.form.get("email", "").strip()

        if not name or not email:

            photo_result = connection.execute(
                """
                SELECT photo_path
                FROM profile_photos
                WHERE user_id = ?
                """,
                (session["user_id"],)
            ).fetchone()

            connection.close()

            return render_template(
                "profile.html",
                user_name=session["user_name"],
                user_email=session["user_email"],
                profile_photo=photo_result["photo_path"]
                if photo_result else None,
                error="Name and email are required."
            )

        try:

            # -----------------------------------------
            # UPDATE NAME AND EMAIL
            # -----------------------------------------

            connection.execute(
                """
                UPDATE users
                SET name = ?, email = ?
                WHERE id = ?
                """,
                (
                    name,
                    email,
                    session["user_id"]
                )
            )


            # -----------------------------------------
            # PROFILE PHOTO UPLOAD
            # -----------------------------------------

            photo = request.files.get("profile_photo")

            if photo and photo.filename:

                allowed_extensions = {
                    "jpg",
                    "jpeg",
                    "png"
                }

                original_name = secure_filename(
                    photo.filename
                )

                extension = (
                    original_name
                    .rsplit(".", 1)[1]
                    .lower()
                    if "." in original_name
                    else ""
                )

                if extension not in allowed_extensions:

                    connection.close()

                    return render_template(
                        "profile.html",
                        user_name=name,
                        user_email=email,
                        profile_photo=None,
                        error="Only JPG, JPEG and PNG images are allowed."
                    )


                # Create upload directory
                upload_folder = os.path.join(
                    "static",
                    "uploads",
                    "profile"
                )

                os.makedirs(
                    upload_folder,
                    exist_ok=True
                )


                # Generate unique filename
                filename = (
                    f"user_{session['user_id']}_"
                    f"{uuid4().hex}.{extension}"
                )


                file_path = os.path.join(
                    upload_folder,
                    filename
                )


                # Save image
                photo.save(file_path)


                # Save relative path in database
                photo_path = (
                    f"uploads/profile/{filename}"
                )


                # Check old photo
                old_photo = connection.execute(
                    """
                    SELECT photo_path
                    FROM profile_photos
                    WHERE user_id = ?
                    """,
                    (session["user_id"],)
                ).fetchone()


                # Update database
                connection.execute(
                    """
                    INSERT INTO profile_photos
                        (user_id, photo_path)
                    VALUES
                        (?, ?)
                    ON CONFLICT(user_id)
                    DO UPDATE SET
                        photo_path = excluded.photo_path
                    """,
                    (
                        session["user_id"],
                        photo_path
                    )
                )


                # Delete old image
                if old_photo and old_photo["photo_path"]:

                    old_file = os.path.join(
                        "static",
                        old_photo["photo_path"]
                    )

                    if (
                        os.path.exists(old_file)
                        and old_file != file_path
                    ):
                        try:
                            os.remove(old_file)
                        except Exception:
                            pass


            # -----------------------------------------
            # SAVE DATABASE CHANGES
            # -----------------------------------------

            connection.commit()


            # Update session
            session["user_name"] = name
            session["user_email"] = email


            connection.close()


            return redirect(
                url_for("dashboard")
            )


        except Exception as error:

            connection.rollback()
            connection.close()

            print(
                "Profile Update Error:",
                error
            )

            return render_template(
                "profile.html",
                user_name=session["user_name"],
                user_email=session["user_email"],
                profile_photo=None,
                error="Unable to update profile."
            )


    # -----------------------------------------
    # GET PROFILE PHOTO
    # -----------------------------------------

    photo_result = connection.execute(
        """
        SELECT photo_path
        FROM profile_photos
        WHERE user_id = ?
        """,
        (session["user_id"],)
    ).fetchone()


    connection.close()


    return render_template(
        "profile.html",
        user_name=session["user_name"],
        user_email=session["user_email"],
        profile_photo=photo_result["photo_path"]
        if photo_result else None
    )    
    
@app.route("/create-resume", methods=["GET", "POST"])
def create_resume():

    if "user_id" not in session:
        return redirect(url_for("home"))

    if request.method == "POST":

        resume_name = request.form["resume_name"].strip()
        resume_type = request.form["resume_type"]
        style = request.form["style"]

        if not resume_name:
            flash("Please enter a resume name.", "error")
            return redirect(url_for("create_resume"))

        connection = get_db()

        cursor = connection.execute(
            """
            INSERT INTO resumes
            (user_id, resume_name, resume_type, style, resume_data)
            VALUES (?, ?, ?, ?, ?)
            """,
            (
                session["user_id"],
                resume_name,
                resume_type,
                style,
                json.dumps({})
            )
        )

        resume_id = cursor.lastrowid

        connection.commit()
        connection.close()

        return redirect(
            url_for(
                "edit_resume",
                resume_id=resume_id
            )
        )

    return render_template("create_resume.html")
@app.route("/resume/<int:resume_id>/delete", methods=["POST"])
def delete_resume(resume_id):

    if "user_id" not in session:
        return redirect(url_for("home"))

    connection = get_db()

    resume = connection.execute(
        """
        SELECT id
        FROM resumes
        WHERE id = ? AND user_id = ?
        """,
        (resume_id, session["user_id"])
    ).fetchone()

    if not resume:
        connection.close()

        flash(
            "Resume not found.",
            "error"
        )

        return redirect(
            url_for("dashboard")
        )

    connection.execute(
        """
        DELETE FROM resumes
        WHERE id = ? AND user_id = ?
        """,
        (resume_id, session["user_id"])
    )

    connection.commit()
    connection.close()

    flash(
        "Resume deleted successfully.",
        "success"
    )

    return redirect(
        url_for("dashboard")
    )
@app.route("/resume/<int:resume_id>/edit", methods=["GET", "POST"])
def edit_resume(resume_id):

    if "user_id" not in session:
        return redirect(url_for("home"))

    connection = get_db()

    resume = connection.execute(
        """
        SELECT *
        FROM resumes
        WHERE id = ? AND user_id = ?
        """,
        (resume_id, session["user_id"])
    ).fetchone()

    if not resume:
        connection.close()
        flash("Resume not found.", "error")
        return redirect(url_for("dashboard"))

    if request.method == "POST":

        # -----------------------------
        # EDUCATION
        # -----------------------------

        education_degrees = request.form.getlist(
            "education_degree[]"
        )

        education_institutions = request.form.getlist(
            "education_institution[]"
        )

        education_starts = request.form.getlist(
            "education_start[]"
        )

        education_ends = request.form.getlist(
            "education_end[]"
        )

        education_scores = request.form.getlist(
            "education_score[]"
        )

        education = []

        for i in range(len(education_degrees)):

            education.append({
                "degree": education_degrees[i],

                "institution": education_institutions[i]
                if i < len(education_institutions)
                else "",

                "start": education_starts[i]
                if i < len(education_starts)
                else "",

                "end": education_ends[i]
                if i < len(education_ends)
                else "",

                "score": education_scores[i]
                if i < len(education_scores)
                else ""
            })


        # -----------------------------
        # PROJECTS
        # -----------------------------

        project_names = request.form.getlist(
            "project_name[]"
        )

        project_technologies = request.form.getlist(
            "project_technologies[]"
        )

        project_descriptions = request.form.getlist(
            "project_description[]"
        )

        projects = []

        for i in range(len(project_names)):

            projects.append({
                "name": project_names[i],

                "technologies": project_technologies[i]
                if i < len(project_technologies)
                else "",

                "description": project_descriptions[i]
                if i < len(project_descriptions)
                else ""
            })


        # -----------------------------
        # INTERNSHIPS
        # -----------------------------

        internship_companies = request.form.getlist(
            "internship_company[]"
        )

        internship_roles = request.form.getlist(
            "internship_role[]"
        )

        internship_starts = request.form.getlist(
            "internship_start[]"
        )

        internship_ends = request.form.getlist(
            "internship_end[]"
        )

        internship_descriptions = request.form.getlist(
            "internship_description[]"
        )

        internships = []

        for i in range(len(internship_companies)):

            internships.append({
                "company": internship_companies[i],

                "role": internship_roles[i]
                if i < len(internship_roles)
                else "",

                "start": internship_starts[i]
                if i < len(internship_starts)
                else "",

                "end": internship_ends[i]
                if i < len(internship_ends)
                else "",

                "description": internship_descriptions[i]
                if i < len(internship_descriptions)
                else ""
            })


        # -----------------------------
        # WORK EXPERIENCE
        # -----------------------------

        experience_companies = request.form.getlist(
            "experience_company[]"
        )

        experience_roles = request.form.getlist(
            "experience_role[]"
        )

        experience_starts = request.form.getlist(
            "experience_start[]"
        )

        experience_ends = request.form.getlist(
            "experience_end[]"
        )

        experience_descriptions = request.form.getlist(
            "experience_description[]"
        )

        experiences = []

        for i in range(len(experience_companies)):

            experiences.append({
                "company": experience_companies[i],

                "role": experience_roles[i]
                if i < len(experience_roles)
                else "",

                "start": experience_starts[i]
                if i < len(experience_starts)
                else "",

                "end": experience_ends[i]
                if i < len(experience_ends)
                else "",

                "description": experience_descriptions[i]
                if i < len(experience_descriptions)
                else ""
            })


        # -----------------------------
        # COMPLETE RESUME DATA
        # -----------------------------

        resume_data = {

            "full_name":
                request.form.get("full_name", "").strip(),

            "professional_title":
                request.form.get("professional_title", "").strip(),

            "email":
                request.form.get("email", "").strip(),

            "phone":
                request.form.get("phone", "").strip(),

            "location":
                request.form.get("location", "").strip(),

            "linkedin":
                request.form.get("linkedin", "").strip(),

            "github":
                request.form.get("github", "").strip(),

            "summary":
                request.form.get("summary", "").strip(),

            "career_objective":
                request.form.get("career_objective", "").strip(),

            "education":
                education,

            "skills":
                request.form.get("skills", "").strip(),

            "projects":
                projects,

            "internship":
                request.form.get("internship", "").strip(),

            "internships":
                internships,

            "work_experience":
                request.form.get("work_experience", "").strip(),

            "experiences":
                experiences,

            "certifications":
                request.form.get("certifications", "").strip(),

            "achievements":
                request.form.get("achievements", "").strip(),

            "languages":
                request.form.get("languages", "").strip(),

            "interests":
                request.form.get("interests", "").strip()
        }


        # -----------------------------
        # SAVE RESUME
        # -----------------------------

        connection.execute(
            """
            UPDATE resumes
            SET resume_data = ?,
                updated_at = CURRENT_TIMESTAMP
            WHERE id = ? AND user_id = ?
            """,
            (
                json.dumps(resume_data),
                resume_id,
                session["user_id"]
            )
        )

        connection.commit()
        connection.close()

        flash(
            "Resume details saved successfully.",
            "success"
        )

        return redirect(
            url_for(
                "edit_resume",
                resume_id=resume_id
            )
        )


    # -----------------------------
    # LOAD SAVED RESUME DATA
    # -----------------------------

    try:

        resume_data = json.loads(
            resume["resume_data"] or "{}"
        )

    except (json.JSONDecodeError, TypeError):

        resume_data = {}


    # Make sure dynamic sections are lists

    if not isinstance(
        resume_data.get("education"),
        list
    ):
        resume_data["education"] = []


    if not isinstance(
        resume_data.get("projects"),
        list
    ):
        resume_data["projects"] = []


    if not isinstance(
        resume_data.get("internships"),
        list
    ):
        resume_data["internships"] = []


    if not isinstance(
        resume_data.get("experiences"),
        list
    ):
        resume_data["experiences"] = []


    connection.close()


    return render_template(
        "resume_builder.html",
        resume=resume,
        resume_data=resume_data
    )
@app.route("/resume/<int:resume_id>/preview")
def preview_resume(resume_id):

    if "user_id" not in session:
        return redirect(url_for("home"))

    connection = get_db()

    resume = connection.execute(
        """
        SELECT *
        FROM resumes
        WHERE id = ? AND user_id = ?
        """,
        (resume_id, session["user_id"])
    ).fetchone()

    connection.close()

    if not resume:
        flash("Resume not found.", "error")
        return redirect(url_for("dashboard"))

    try:
        resume_data = json.loads(
            resume["resume_data"] or "{}"
        )
    except (json.JSONDecodeError, TypeError):
        resume_data = {}

    return render_template(
        "resume_preview.html",
        resume=resume,
        resume_data=resume_data
    )

@app.route("/resume/<int:resume_id>/download-pdf")
def download_pdf(resume_id):

    if "user_id" not in session:
        return redirect(url_for("home"))

    connection = get_db()

    resume = connection.execute(
        """
        SELECT *
        FROM resumes
        WHERE id = ? AND user_id = ?
        """,
        (resume_id, session["user_id"])
    ).fetchone()

    connection.close()

    if not resume:
        flash("Resume not found.", "error")
        return redirect(url_for("dashboard"))

    try:
        resume_data = json.loads(
            resume["resume_data"] or "{}"
        )
    except (json.JSONDecodeError, TypeError):
        resume_data = {}

    html = render_template(
        "resume_pdf.html",
        resume=resume,
        resume_data=resume_data
    )

    pdf_file = BytesIO()

    HTML(
        string=html,
        base_url=request.url_root
    ).write_pdf(pdf_file)

    pdf_file.seek(0)

    response = make_response(pdf_file.read())

    response.headers["Content-Type"] = "application/pdf"

    response.headers["Content-Disposition"] = (
        f'attachment; filename="{resume["resume_name"]}.pdf"'
    )

    return response
@app.route("/resume/<int:resume_id>/download-word")
def download_word(resume_id):

    if "user_id" not in session:
        return redirect(url_for("home"))

    connection = get_db()

    resume = connection.execute(
        """
        SELECT *
        FROM resumes
        WHERE id = ? AND user_id = ?
        """,
        (resume_id, session["user_id"])
    ).fetchone()

    connection.close()

    if not resume:
        flash("Resume not found.", "error")
        return redirect(url_for("dashboard"))

    try:
        resume_data = json.loads(
            resume["resume_data"] or "{}"
        )
    except (json.JSONDecodeError, TypeError):
        resume_data = {}

    # ---------------------------------
    # CREATE WORD DOCUMENT
    # ---------------------------------

    document = Document()

    # Default font
    styles = document.styles

    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)

    # ---------------------------------
    # NAME
    # ---------------------------------

    name = resume_data.get("full_name", "")

    if name:
        paragraph = document.add_paragraph()

        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run(name)

        run.bold = True
        run.font.size = Pt(20)

    # ---------------------------------
    # PROFESSIONAL TITLE
    # ---------------------------------

    professional_title = resume_data.get(
        "professional_title",
        ""
    )

    if professional_title:

        paragraph = document.add_paragraph()

        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run(
            professional_title
        )

        run.bold = True
        run.font.size = Pt(11)

    # ---------------------------------
    # CONTACT INFORMATION
    # ---------------------------------

    contact = []

    if resume_data.get("email"):
        contact.append(
            resume_data.get("email")
        )

    if resume_data.get("phone"):
        contact.append(
            resume_data.get("phone")
        )

    if resume_data.get("location"):
        contact.append(
            resume_data.get("location")
        )

    if contact:

        paragraph = document.add_paragraph()

        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        paragraph.add_run(
            " | ".join(contact)
        )

    # ---------------------------------
    # LINKS
    # ---------------------------------

    links = []

    if resume_data.get("linkedin"):
        links.append(
            "LinkedIn: " +
            resume_data.get("linkedin")
        )

    if resume_data.get("github"):
        links.append(
            "GitHub: " +
            resume_data.get("github")
        )

    if links:

        paragraph = document.add_paragraph()

        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        paragraph.add_run(
            " | ".join(links)
        )

    # ---------------------------------
    # HELPER FUNCTION
    # ---------------------------------

    def add_section_title(title):

        paragraph = document.add_paragraph()

        run = paragraph.add_run(title)

        run.bold = True
        run.font.size = Pt(12)

    # ---------------------------------
    # SUMMARY
    # ---------------------------------

    if resume_data.get("summary"):

        add_section_title(
            "PROFESSIONAL SUMMARY"
        )

        document.add_paragraph(
            resume_data.get("summary")
        )

    # ---------------------------------
    # CAREER OBJECTIVE
    # ---------------------------------

    if resume_data.get("career_objective"):

        add_section_title(
            "CAREER OBJECTIVE"
        )

        document.add_paragraph(
            resume_data.get("career_objective")
        )

    # ---------------------------------
    # EDUCATION
    # ---------------------------------

    education = resume_data.get(
        "education",
        []
    )

    if education:

        add_section_title("EDUCATION")

        for item in education:

            paragraph = document.add_paragraph()

            degree = item.get(
                "degree",
                ""
            )

            institution = item.get(
                "institution",
                ""
            )

            start = item.get(
                "start",
                ""
            )

            end = item.get(
                "end",
                ""
            )

            score = item.get(
                "score",
                ""
            )

            text = degree

            if institution:
                text += " - " + institution

            dates = ""

            if start:
                dates += start

            if end:
                if dates:
                    dates += " - "
                dates += end

            if dates:
                text += " (" + dates + ")"

            if score:
                text += " | " + score

            paragraph.add_run(
                text
            )

    # ---------------------------------
    # SKILLS
    # ---------------------------------

    if resume_data.get("skills"):

        add_section_title("SKILLS")

        document.add_paragraph(
            resume_data.get("skills")
        )

    # ---------------------------------
    # PROJECTS
    # ---------------------------------

    projects = resume_data.get(
        "projects",
        []
    )

    if projects:

        add_section_title("PROJECTS")

        for project in projects:

            name = project.get(
                "name",
                ""
            )

            technologies = project.get(
                "technologies",
                ""
            )

            description = project.get(
                "description",
                ""
            )

            if name:

                paragraph = document.add_paragraph()

                run = paragraph.add_run(
                    name
                )

                run.bold = True

            if technologies:

                document.add_paragraph(
                    "Technologies: " +
                    technologies
                )

            if description:

                document.add_paragraph(
                    description
                )

    # ---------------------------------
    # INTERNSHIPS
    # ---------------------------------

    internships = resume_data.get(
        "internships",
        []
    )

    if internships:

        add_section_title(
            "INTERNSHIP"
        )

        for internship in internships:

            role = internship.get(
                "role",
                ""
            )

            company = internship.get(
                "company",
                ""
            )

            start = internship.get(
                "start",
                ""
            )

            end = internship.get(
                "end",
                ""
            )

            description = internship.get(
                "description",
                ""
            )

            paragraph = document.add_paragraph()

            text = ""

            if role:
                text += role

            if company:
                if text:
                    text += " - "
                text += company

            if start or end:

                if text:
                    text += " | "

                text += start

                if start and end:
                    text += " - "

                text += end

            run = paragraph.add_run(
                text
            )

            run.bold = True

            if description:

                document.add_paragraph(
                    description
                )

    # ---------------------------------
    # WORK EXPERIENCE
    # ---------------------------------

    experiences = resume_data.get(
        "experiences",
        []
    )

    if experiences:

        add_section_title(
            "WORK EXPERIENCE"
        )

        for experience in experiences:

            role = experience.get(
                "role",
                ""
            )

            company = experience.get(
                "company",
                ""
            )

            start = experience.get(
                "start",
                ""
            )

            end = experience.get(
                "end",
                ""
            )

            description = experience.get(
                "description",
                ""
            )

            paragraph = document.add_paragraph()

            text = ""

            if role:
                text += role

            if company:
                if text:
                    text += " - "
                text += company

            if start or end:

                if text:
                    text += " | "

                text += start

                if start and end:
                    text += " - "

                text += end

            run = paragraph.add_run(
                text
            )

            run.bold = True

            if description:

                document.add_paragraph(
                    description
                )

    # ---------------------------------
    # CERTIFICATIONS
    # ---------------------------------

    if resume_data.get(
        "certifications"
    ):

        add_section_title(
            "CERTIFICATIONS"
        )

        document.add_paragraph(
            resume_data.get(
                "certifications"
            )
        )

    # ---------------------------------
    # ACHIEVEMENTS
    # ---------------------------------

    if resume_data.get(
        "achievements"
    ):

        add_section_title(
            "ACHIEVEMENTS"
        )

        document.add_paragraph(
            resume_data.get(
                "achievements"
            )
        )

    # ---------------------------------
    # LANGUAGES
    # ---------------------------------

    if resume_data.get(
        "languages"
    ):

        add_section_title(
            "LANGUAGES"
        )

        document.add_paragraph(
            resume_data.get(
                "languages"
            )
        )

    # ---------------------------------
    # INTERESTS
    # ---------------------------------

    if resume_data.get(
        "interests"
    ):

        add_section_title(
            "INTERESTS"
        )

        document.add_paragraph(
            resume_data.get(
                "interests"
            )
        )

    # ---------------------------------
    # SAVE WORD FILE TO MEMORY
    # ---------------------------------

    word_file = BytesIO()

    document.save(word_file)

    word_file.seek(0)

    response = make_response(
        word_file.read()
    )

    response.headers[
        "Content-Type"
    ] = (
        "application/vnd.openxmlformats-officedocument."
        "wordprocessingml.document"
    )

    response.headers[
        "Content-Disposition"
    ] = (
        f'attachment; filename="{resume["resume_name"]}.docx"'
    )

    return response
@app.route("/resume/<int:resume_id>/download-image")
def download_image(resume_id):

    if "user_id" not in session:
        return redirect(url_for("home"))

    connection = get_db()

    resume = connection.execute(
        """
        SELECT *
        FROM resumes
        WHERE id = ? AND user_id = ?
        """,
        (resume_id, session["user_id"])
    ).fetchone()

    connection.close()

    if not resume:
        flash("Resume not found.", "error")
        return redirect(url_for("dashboard"))

    try:
        resume_data = json.loads(
            resume["resume_data"] or "{}"
        )
    except (json.JSONDecodeError, TypeError):
        resume_data = {}

    # Render the resume as HTML
    html = render_template(
        "resume_pdf.html",
        resume=resume,
        resume_data=resume_data
    )

    # Convert HTML to PDF in memory
    pdf_file = BytesIO()

    HTML(
        string=html,
        base_url=request.url_root
    ).write_pdf(pdf_file)

    pdf_file.seek(0)

    # Open PDF using PyMuPDF
    pdf_document = fitz.open(
        stream=pdf_file.read(),
        filetype="pdf"
    )

    if pdf_document.page_count == 0:
        pdf_document.close()

        flash(
            "Unable to create resume image.",
            "error"
        )

        return redirect(
            url_for(
                "preview_resume",
                resume_id=resume_id
            )
        )

    # Get first A4 page
    page = pdf_document[0]

    # Render page as high-resolution PNG
    matrix = fitz.Matrix(2.5, 2.5)

    pixmap = page.get_pixmap(
        matrix=matrix,
        alpha=False
    )

    image_data = pixmap.tobytes(
        "png"
    )

    pdf_document.close()

    # Send PNG to browser
    response = make_response(
        image_data
    )

    response.headers[
        "Content-Type"
    ] = "image/png"

    response.headers[
        "Content-Disposition"
    ] = (
        f'attachment; filename="{resume["resume_name"]}.png"'
    )

    return response       

# -----------------------------
# LOGOUT
# -----------------------------

@app.route("/logout")
def logout():

    session.clear()

    return redirect(url_for("home"))


# -----------------------------
# START APPLICATION
# -----------------------------
# Initialize database when the app starts
init_db()

# Local development
if __name__ == "__main__":
    app.run(
        host="0.0.0.0",
        port=5000,
        debug=True
    )